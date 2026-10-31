"""Owned file library storage and simple text templates."""

import io
import os
import re
import uuid
import zipfile
from pathlib import Path
from typing import Annotated

from fastapi import APIRouter, Depends, HTTPException, Request, UploadFile
from fastapi.responses import FileResponse
from pydantic import BaseModel, ConfigDict, Field

from business_models import BusinessInput, Name, ShortText
from database import connection
from tenancy import TenantContext, require_tenant

router=APIRouter(prefix='/api',tags=['documents and templates'])
MAX_FILE_SIZE=12*1024*1024
ALLOWED={'.txt','.md','.csv','.pdf','.docx','.xlsx','.png','.jpg','.jpeg','.webp'}


class LibraryInput(BusinessInput):
    name:Name
    description:ShortText=''
    root_path:ShortText=''


class TemplateInput(BusinessInput):
    name:Name
    content:str=Field(max_length=50000)


class TemplateValues(BaseModel):
    model_config=ConfigDict(extra='forbid')
    customer:str=Field('',max_length=2000)
    company:str=Field('',max_length=2000)
    products:str=Field('',max_length=10000)
    contact:str=Field('',max_length=2000)


def document_context(request:Request,tenant:TenantContext=Depends(require_tenant)):
    return request.app.state.db_path,tenant


Context=Annotated[tuple,Depends(document_context)]


def library(path,wid,lid,company_id=None):
    with connection(path) as conn:
        if company_id is None:
            row=conn.execute('SELECT * FROM libraries WHERE id=? AND workspace_id=?',(lid,wid)).fetchone()
        else:
            row=conn.execute('SELECT * FROM libraries WHERE id=? AND workspace_id=? AND company_id=?',(lid,wid,company_id)).fetchone()
    if row is None:raise HTTPException(404,'Library not found')
    return dict(row)


def validated_folder(raw: str) -> Path:
    if not raw or not raw.strip():
        return Path()
    root = Path(raw.strip()).expanduser().resolve()
    home = Path.home().resolve()
    if not root.is_dir():
        raise HTTPException(422, '文件夹不存在或不是目录。')
    if root == home or root.parent == root:
        raise HTTPException(422, '不能直接关联家目录或磁盘根目录。')
    forbidden_roots = [
        home/'.ssh', home/'.trade', home/'.hermes', home/'.codex',
        Path(os.environ.get('SystemRoot', 'C:/Windows')),
        Path(os.environ.get('ProgramFiles', 'C:/Program Files')),
        Path(os.environ.get('ProgramData', 'C:/ProgramData')),
    ]
    for forbidden in forbidden_roots:
        try:
            root.relative_to(forbidden.resolve())
        except ValueError:
            continue
        raise HTTPException(422, '不能关联系统或应用敏感目录。')
    return root


def document(path,wid,did,company_id=None):
    with connection(path) as conn:
        row=conn.execute('''
            SELECT documents.* FROM documents JOIN libraries
              ON libraries.id=documents.library_id AND libraries.workspace_id=documents.workspace_id
            WHERE documents.id=? AND documents.workspace_id=? AND documents.deleted=0
              AND (? IS NULL OR libraries.company_id=?)
        ''',(did,wid,company_id,company_id)).fetchone()
    if row is None:raise HTTPException(404,'Document not found')
    return dict(row)


def public_document(row):
    return {key:row[key] for key in ['id','library_id','filename','size','extraction_note','created_at']}


def file_path(path,wid,storage_name):
    root=(Path(path).parent/'workspaces'/str(wid)/'documents').resolve()
    target=(root/storage_name).resolve()
    if target.parent!=root:raise HTTPException(404,'Document not found')
    return target


def extract_text(data,suffix):
    if suffix in {'.txt','.md','.csv'}:
        try:return data.decode('utf-8-sig')[:100000],'已读取 UTF-8 文本（最多 100000 字符）'
        except UnicodeDecodeError:return '', '不是 UTF-8 文本；原文件已保存，未提取文字。'
    if suffix in {'.png','.jpg','.jpeg','.webp'}:return '', '图片/技术图纸已保存；本阶段不执行视觉识别。'
    try:
        if suffix in {'.docx','.xlsx'}:
            with zipfile.ZipFile(io.BytesIO(data)) as archive:
                if len(archive.infolist())>1000 or sum(item.file_size for item in archive.infolist())>40*1024*1024:
                    return '', '压缩内容超出解析限额；原文件已保存。'
            if suffix=='.docx':
                from docx import Document
                doc=Document(io.BytesIO(data))
                text='\n'.join(paragraph.text for paragraph in doc.paragraphs)
                for table in doc.tables:
                    text+='\n'+'\n'.join('\t'.join(cell.text for cell in row.cells) for row in table.rows)
            else:
                from openpyxl import load_workbook
                book=load_workbook(io.BytesIO(data),read_only=True,data_only=True)
                try:
                    parts=[]
                    for sheet in book.worksheets[:10]:
                        parts.append(sheet.title)
                        for row in sheet.iter_rows(max_row=500,max_col=30,values_only=True):
                            parts.append('\t'.join(str(value) if value is not None else '' for value in row))
                    text='\n'.join(parts)
                finally:book.close()
        elif suffix=='.pdf':
            import pypdfium2 as pdfium
            pdf=pdfium.PdfDocument(data)
            try:
                parts=[]
                for index in range(min(len(pdf),50)):
                    page=pdf[index]
                    try:
                        text_page=page.get_textpage()
                        try:parts.append(text_page.get_text_range()[:10000])
                        finally:text_page.close()
                    finally:page.close()
                text='\n'.join(parts)
            finally:pdf.close()
        else:return '', '原文件已保存，未提取文字。'
        return text[:100000], ('已提取文本；PDF最多50页，表格最多10个工作表。' if text.strip() else '未提取到文字，可能是扫描件或图纸；未执行 OCR。')
    except ImportError:return '', '本地缺少此格式的解析依赖；原文件已保存。'
    except Exception:return '', '此文件暂时无法提取文字；原文件已保存。'


def scan_folder(root: Path, max_files: int = 500) -> list[Path]:
    files = []
    for candidate in root.rglob('*'):
        if len(files) >= max_files:
            break
        try:
            relative = candidate.relative_to(root)
        except ValueError:
            continue
        if len(relative.parts) > 4 or candidate.is_symlink() or not candidate.is_file():
            continue
        if candidate.suffix.lower() not in ALLOWED:
            continue
        try:
            if candidate.stat().st_size > MAX_FILE_SIZE:
                continue
        except OSError:
            continue
        files.append(candidate)
    return files


@router.get('/libraries')
def list_libraries(context:Context):
    path,tenant=context
    with connection(path) as conn:
        rows=conn.execute('SELECT libraries.*,(SELECT count(*) FROM documents WHERE library_id=libraries.id AND workspace_id=libraries.workspace_id AND deleted=0) AS file_count FROM libraries WHERE workspace_id=? AND company_id=? ORDER BY id DESC',(tenant.workspace_id,tenant.company_id)).fetchall()
    return [dict(row) for row in rows]


@router.post('/libraries',status_code=201)
def create_library(payload:LibraryInput,context:Context):
    path,tenant=context
    root = validated_folder(payload.root_path)
    root_path = str(root) if str(root) != '.' else ''
    with connection(path) as conn:
        cursor=conn.execute('INSERT INTO libraries(workspace_id,company_id,name,description,root_path) VALUES (?,?,?,?,?)',(tenant.workspace_id,tenant.company_id,payload.name,payload.description,root_path))
    return library(path,tenant.workspace_id,cursor.lastrowid,tenant.company_id)


@router.post('/libraries/{lid}/scan')
async def scan_library(lid:int,context:Context):
    path,tenant=context;wid=tenant.workspace_id;row=library(path,wid,lid,tenant.company_id)
    if not row['root_path']:
        raise HTTPException(422, '这个文档库没有关联本地文件夹。')
    root=validated_folder(row['root_path'])
    from starlette.concurrency import run_in_threadpool

    def scan_and_extract():
        result=[]
        for candidate in scan_folder(root):
            try:
                data=candidate.read_bytes()
            except OSError:
                continue
            text,note=extract_text(data,candidate.suffix.lower())
            result.append((candidate,data,len(data),text,note))
        return result
    entries=await run_in_threadpool(scan_and_extract)
    with connection(path) as conn:
        for candidate,data,size,text,note in entries:
            existing=conn.execute(
                'SELECT id FROM documents WHERE workspace_id=? AND library_id=? AND source_path=?',
                (wid,lid,str(candidate)),
            ).fetchone()
            if existing:
                conn.execute(
                    'UPDATE documents SET filename=?,size=?,extracted_text=?,extraction_note=?,deleted=0 WHERE id=? AND workspace_id=?',
                    (candidate.name,size,text,note,existing['id'],wid),
                )
            else:
                conn.execute(
                    'INSERT INTO documents(workspace_id,library_id,filename,storage_name,source_path,size,extracted_text,extraction_note) VALUES (?,?,?,?,?,?,?,?)',
                    (wid,lid,candidate.name,'',str(candidate),size,text,note),
                )
    return {'found':len(entries)}


@router.get('/libraries/{lid}/files')
def list_files(lid:int,context:Context):
    path,tenant=context
    library(path,tenant.workspace_id,lid,tenant.company_id)
    with connection(path) as conn:
        rows=conn.execute('SELECT * FROM documents WHERE workspace_id=? AND library_id=? AND deleted=0 ORDER BY id DESC',(tenant.workspace_id,lid)).fetchall()
    return [public_document(row) for row in rows]


@router.post('/libraries/{lid}/files',status_code=201)
async def upload_document(lid:int,file:UploadFile,context:Context):
    path,tenant=context;wid=tenant.workspace_id
    library(path,wid,lid,tenant.company_id)
    filename=(file.filename or 'document').replace('\\','/').rsplit('/',1)[-1]
    filename=re.sub(r'[\x00-\x1f\x7f]','',filename)[:200]
    suffix=Path(filename).suffix.lower()
    if suffix not in ALLOWED:raise HTTPException(422,'支持 txt/md/csv/pdf/docx/xlsx 和常见图片格式。')
    data=bytearray()
    try:
        while chunk:=await file.read(65536):
            data.extend(chunk)
            if len(data)>MAX_FILE_SIZE:raise HTTPException(413,'文件不能超过 12 MB。')
    finally:await file.close()
    if not data:raise HTTPException(422,'文件为空。')
    # Only generated server-side names are used on disk.
    storage_name=uuid.uuid4().hex+suffix
    target=file_path(path,wid,storage_name)
    target.parent.mkdir(parents=True,exist_ok=True)
    target.write_bytes(data)
    from starlette.concurrency import run_in_threadpool
    text,note=await run_in_threadpool(extract_text,bytes(data),suffix)
    try:
        with connection(path) as conn:
            cursor=conn.execute('INSERT INTO documents(workspace_id,library_id,filename,storage_name,size,extracted_text,extraction_note) VALUES (?,?,?,?,?,?,?)',(wid,lid,filename,storage_name,len(data),text,note))
    except Exception:
        target.unlink(missing_ok=True)
        raise
    return public_document(document(path,wid,cursor.lastrowid,tenant.company_id))


@router.get('/documents/{did}/download')
def download_document(did:int,context:Context):
    path,tenant=context
    doc=document(path,tenant.workspace_id,did,tenant.company_id)
    if doc.get('source_path'):
        row=library(path,tenant.workspace_id,doc['library_id'],tenant.company_id)
        root=validated_folder(row['root_path'])
        target=Path(doc['source_path']).resolve()
        try:
            target.relative_to(root)
        except ValueError:
            raise HTTPException(404,'Document file not found')
    else:
        target=file_path(path,tenant.workspace_id,doc['storage_name'])
    if not target.is_file():raise HTTPException(404,'Document file not found')
    return FileResponse(target,filename=doc['filename'],media_type='application/octet-stream',headers={'X-Content-Type-Options':'nosniff','Cache-Control':'no-store'})


@router.get('/documents/{did}/text')
def document_text(did:int,context:Context):
    path,tenant=context;doc=document(path,tenant.workspace_id,did,tenant.company_id)
    return {'id':did,'filename':doc['filename'],'text':doc['extracted_text'],'note':doc['extraction_note']}


@router.delete('/documents/{did}')
def delete_document(did:int,context:Context):
    path,tenant=context
    document(path,tenant.workspace_id,did,tenant.company_id)
    with connection(path) as conn:conn.execute('UPDATE documents SET deleted=1 WHERE id=? AND workspace_id=?',(did,tenant.workspace_id))
    return {'ok':True}


@router.get('/templates')
def list_templates(context:Context):
    path,tenant=context
    with connection(path) as conn:rows=conn.execute('SELECT id,name,content,updated_at FROM templates WHERE workspace_id=? ORDER BY id DESC',(tenant.workspace_id,)).fetchall()
    return [dict(row) for row in rows]


@router.post('/templates',status_code=201)
def create_template(payload:TemplateInput,context:Context):
    path,tenant=context
    with connection(path) as conn:cursor=conn.execute('INSERT INTO templates(workspace_id,name,content) VALUES (?,?,?)',(tenant.workspace_id,payload.name,payload.content))
    return {'id':cursor.lastrowid,**payload.model_dump()}


def template(path,wid,tid):
    with connection(path) as conn:row=conn.execute('SELECT id,name,content FROM templates WHERE id=? AND workspace_id=?',(tid,wid)).fetchone()
    if row is None:raise HTTPException(404,'Template not found')
    return dict(row)


@router.put('/templates/{tid}')
def update_template(tid:int,payload:TemplateInput,context:Context):
    path,tenant=context;template(path,tenant.workspace_id,tid)
    with connection(path) as conn:conn.execute("UPDATE templates SET name=?,content=?,updated_at=strftime('%Y-%m-%dT%H:%M:%fZ','now') WHERE id=? AND workspace_id=?",(payload.name,payload.content,tid,tenant.workspace_id))
    return {'id':tid,**payload.model_dump()}


@router.delete('/templates/{tid}')
def delete_template(tid:int,context:Context):
    path,tenant=context;template(path,tenant.workspace_id,tid)
    with connection(path) as conn:conn.execute('DELETE FROM templates WHERE id=? AND workspace_id=?',(tid,tenant.workspace_id))
    return {'ok':True}


@router.post('/templates/{tid}/render')
def render_template(tid:int,payload:TemplateValues,context:Context):
    path,tenant=context;row=template(path,tenant.workspace_id,tid);values=payload.model_dump()
    # One plain-text substitution pass; never evaluate expressions or code.
    content=re.sub(r'\{\{\s*(company|customer|products|contact)\s*\}\}',lambda match:values[match.group(1)],row['content'])
    return {'content':content}
